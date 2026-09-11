"""Cockpit xlsx: visão gerada, nunca fonte. Os testes cobrem três coisas que dão número errado
em silêncio numa planilha — faixa de fórmula inválida, denominador incompleto e afirmação de
rentabilidade que ignora câmbio — mais o contrato de saída do CLI."""
import datetime
import os
import subprocess
import sys
from pathlib import Path

import pytest

from po.carteira import valorar
from test_validar_dados import _anexa, copia_exemplo

pytest.importorskip("openpyxl")

import openpyxl  # noqa: E402

from po.cockpit import gerar_cockpit  # noqa: E402

RAIZ = Path(__file__).resolve().parent.parent
CLI = RAIZ / "scripts" / "gerar_cockpit.py"
AGORA = datetime.datetime(2026, 9, 10, 18, 0)


def _abre(ws_raiz, agora=AGORA):
    caminho = gerar_cockpit(ws_raiz, agora=agora)
    return caminho, openpyxl.load_workbook(caminho)


def _coluna(ws, col, ini=2):
    return [ws.cell(row=r, column=col).value for r in range(ini, ws.max_row + 1)]


def _leiame(wb):
    return [c[0].value or "" for c in wb["LEIAME"].iter_rows(min_col=1, max_col=1)]


def test_gera_as_quatro_abas_e_congela_cabecalho(tmp_path):
    _, wb = _abre(copia_exemplo(tmp_path))
    assert wb.sheetnames == ["LEIAME", "Posições", "Blocos", "Aporte"]
    assert wb["Posições"].freeze_panes == "A2" and wb["Blocos"].freeze_panes == "A2"


def test_numeros_vem_da_valoracao_e_nao_de_recalculo(tmp_path):
    """O cockpit e o ESTADO.md mostram o mesmo número porque os dois chamam `valorar`. Se a
    planilha recalculasse por conta própria, a divergência apareceria primeiro aqui."""
    ws_raiz = copia_exemplo(tmp_path)
    c = valorar(ws_raiz, hoje=AGORA.date())
    _, wb = _abre(ws_raiz)
    pos = wb["Posições"]
    esperado = {(l.ticker, l.qty, l.preco, l.data_cotacao, l.fonte) for l in c.linhas}
    lido = {(pos.cell(row=r, column=1).value, pos.cell(row=r, column=4).value,
             pos.cell(row=r, column=7).value, pos.cell(row=r, column=8).value,
             pos.cell(row=r, column=9).value) for r in range(2, 2 + len(c.linhas))}
    assert lido == esperado


def test_valor_brl_e_formula_sobre_as_proprias_colunas(tmp_path):
    """Valor BRL não é número gravado: é qty × cotação × câmbio, visível na barra de fórmulas.
    Assim o usuário confere a conta sem abrir o CSV."""
    _, wb = _abre(copia_exemplo(tmp_path))
    pos = wb["Posições"]
    assert pos["K2"].value == "=D2*G2*J2"
    assert pos["K4"].value == "=SUM(K2:K3)"       # linha de Total, com 2 posições
    assert pos["N2"].value == "=IF($K$4=0,0,K2/$K$4)"


def test_workspace_vazio_gera_planilha_valida_sem_faixa_invertida(tmp_path):
    """Primeiro arquivo que um usuário novo abre. Com zero posições, uma faixa ingênua viraria
    K2:K1 e o Excel reclamaria do arquivo inteiro."""
    from criar_workspace import criar

    caminho, wb = _abre(criar(tmp_path / "ws", com_git=False, data="2026-09-08"))
    assert caminho.exists()
    formulas = [c.value for linha in wb["Posições"].iter_rows() for c in linha
                if isinstance(c.value, str) and c.value.startswith("=")]
    assert not any(":" in f and _faixa_invertida(f) for f in formulas), formulas
    assert wb["Posições"]["K2"].value == 0        # total literal, não SUM de faixa vazia
    assert "Sem bandas declaradas" in str(wb["Aporte"]["A5"].value)


def _faixa_invertida(formula: str) -> bool:
    import re
    for a, b in re.findall(r"\$?[A-Z]+\$?(\d+):\$?[A-Z]+\$?(\d+)", formula):
        if int(b) < int(a):
            return True
    return False


def test_classe_sem_banda_ganha_linha_e_o_denominador_fica_completo(tmp_path):
    """Se a posição de classe sem banda ficasse fora da tabela Blocos, o Total dela seria menor
    que o das Posições e TODO percentual da carteira sairia errado, em silêncio."""
    ws_raiz = copia_exemplo(tmp_path)
    _anexa(ws_raiz, "dados/posicoes.csv", "BTC,cripto,corretora-br,1,100000.00,BRL")
    _anexa(ws_raiz, "dados/fills.csv", "2026-08-01,BTC,saldo-inicial,1,100000.00,0,corretora-br,BRL")
    _anexa(ws_raiz, "dados/cotacoes.csv", "2026-09-08,18:00,BTC,120000.00,BRL,manual")
    _, wb = _abre(ws_raiz)
    blocos = wb["Blocos"]
    nomes = [v for v in _coluna(blocos, 1) if v]
    assert "cripto" in nomes and nomes[-1] == "Total"
    linha_cripto = nomes.index("cripto") + 2
    assert blocos.cell(row=linha_cripto, column=4).value is None      # sem banda declarada
    assert "sem banda" in blocos.cell(row=linha_cripto, column=7).value
    # o Total soma TODAS as linhas de bloco, cripto inclusive
    assert blocos.cell(row=len(nomes) + 1, column=2).value == f"=SUM(B2:B{len(nomes)})"


def test_aviso_de_cotacao_velha_chega_ao_leiame_e_a_aba_de_aporte(tmp_path):
    """Tela de aporte construída sobre preço velho é o número plausível e errado que esta fase
    existe para caçar: o aviso da valoração não pode morrer no caminho."""
    ws_raiz = copia_exemplo(tmp_path)
    _, wb = _abre(ws_raiz, agora=datetime.datetime(2026, 10, 8, 18, 0))
    texto = _leiame(wb)
    assert any(l.startswith("AVISOS (") for l in texto)
    assert any("mais de 7 dias" in l and "PETR4" in l for l in texto)
    assert "ATENÇÃO" in str(wb["Aporte"]["A1"].value)


def test_sem_aviso_a_aba_de_aporte_nao_grita(tmp_path):
    """Contraprova: aviso que aparece sempre treina o usuário a ignorá-lo."""
    _, wb = _abre(copia_exemplo(tmp_path))
    assert wb["Aporte"]["A1"].value is None
    assert not any(l.startswith("AVISOS (") for l in _leiame(wb))


def test_resultado_e_na_moeda_do_ativo_e_o_leiame_diz_o_que_o_custo_nao_e(tmp_path):
    """`custo_brl` remarca o PM histórico pelo câmbio de hoje, então resultado em BRL mentiria
    para posição em moeda estrangeira. A planilha mostra o ganho na moeda do ativo e declara o
    que o custo é — em vez de exibir um número que ninguém consegue interpretar."""
    _, wb = _abre(copia_exemplo(tmp_path))
    titulos = [c.value for c in wb["Posições"][1]]
    assert "Result. moeda ativo" in titulos
    assert not any(t and "Result" in t and "BRL" in t for t in titulos)
    assert wb["Posições"]["M2"].value == '=IF(E2=0,"",G2/E2-1)'      # cotação ÷ PM, sem câmbio
    texto = " ".join(_leiame(wb))
    assert "câmbio de HOJE" in texto and "não é o valor em reais que saiu da sua conta" in texto


def test_fila_de_aporte_zera_em_vez_de_ficar_negativa(tmp_path):
    """O Sugerido desconta os gaps de quem está acima na fila; quando o aporte acaba a conta fica
    negativa e o MAX(0,...) zera. Sem ele, um bloco de rank pior receberia valor negativo."""
    _, wb = _abre(copia_exemplo(tmp_path))
    ap = wb["Aporte"]
    assert ap["F5"].value.startswith('=IF(E5="","",MAX(0,MIN(D5,$B$2-SUMIFS(')
    assert ap["B2"].value == 0 and ap["C2"].value == "← única célula editável"
    ultimo = ap.max_row
    assert ap.cell(row=ultimo, column=1).value == "Sobra do aporte"
    assert ap.cell(row=ultimo, column=6).value == f"=$B$2-F{ultimo - 1}"


def test_regerar_sobrescreve_e_nao_acumula(tmp_path):
    ws_raiz = copia_exemplo(tmp_path)
    primeiro, _ = _abre(ws_raiz)
    segundo, wb = _abre(ws_raiz, agora=datetime.datetime(2026, 9, 11, 9, 30))
    assert primeiro == segundo
    assert len(list(primeiro.parent.glob("*.xlsx"))) == 1
    assert "11/09/2026 09:30" in _leiame(wb)[1]


def test_dados_que_nao_sustentam_o_numero_nao_geram_planilha(tmp_path):
    ws_raiz = copia_exemplo(tmp_path)
    _anexa(ws_raiz, "dados/posicoes.csv", "VALE3,acoes-br,corretora-br,10,60.00,BRL")
    _anexa(ws_raiz, "dados/fills.csv", "2026-08-01,VALE3,saldo-inicial,10,60.00,0,corretora-br,BRL")
    with pytest.raises(ValueError, match="VALE3 sem cotação"):
        gerar_cockpit(ws_raiz, agora=AGORA)


def _roda(*args, env=None):
    return subprocess.run([sys.executable, str(CLI), *args], capture_output=True, text=True,
                          encoding="utf-8", errors="replace", env=env)


@pytest.mark.slow
def test_cli_gera_e_sai_0(tmp_path):
    ws_raiz = copia_exemplo(tmp_path)
    r = _roda(str(ws_raiz))
    assert r.returncode == 0, r.stdout + r.stderr
    assert "Cockpit gerado:" in r.stdout and "Aporte!B2" in r.stdout
    assert (ws_raiz / "planilhas" / "cockpit.xlsx").exists()


@pytest.mark.slow
def test_cli_raiz_inexistente_sai_1_sem_traceback(tmp_path):
    r = _roda(str(tmp_path / "nao-existe"))
    assert r.returncode == 1
    assert "não existe" in r.stdout and "Traceback" not in r.stderr


@pytest.mark.slow
def test_cli_dados_sujos_sai_1_com_frase(tmp_path):
    ws_raiz = copia_exemplo(tmp_path)
    _anexa(ws_raiz, "dados/posicoes.csv", "VALE3,acoes-br,corretora-br,10,60.00,BRL")
    _anexa(ws_raiz, "dados/fills.csv", "2026-08-01,VALE3,saldo-inicial,10,60.00,0,corretora-br,BRL")
    r = _roda(str(ws_raiz))
    assert r.returncode == 1
    assert "VALE3 sem cotação" in r.stdout and "Traceback" not in r.stderr


@pytest.mark.slow
def test_cli_uso_invalido_sai_2():
    r = _roda()
    assert r.returncode == 2


@pytest.mark.slow
def test_cli_diretorio_no_lugar_do_arquivo_vira_frase(tmp_path):
    """No Windows isso dá PermissionError e em Linux IsADirectoryError: pegar só a primeira
    deixava a outra virar traceback num projeto que roda nos dois."""
    ws_raiz = copia_exemplo(tmp_path)
    alvo = ws_raiz / "planilhas" / "cockpit.xlsx"
    alvo.parent.mkdir(parents=True, exist_ok=True)
    alvo.mkdir()
    r = _roda(str(ws_raiz))
    assert r.returncode == 1
    assert "não consegui ler/gravar" in r.stdout and "Traceback" not in r.stderr


@pytest.mark.slow
def test_cli_sobrevive_a_console_cp1252(tmp_path):
    """O output tem × e ç, que cp1252 até aceita, mas a frase de erro tem ≤ e →."""
    ws_raiz = copia_exemplo(tmp_path)
    env = dict(os.environ, PYTHONIOENCODING="cp1252")
    r = _roda(str(ws_raiz), env=env)
    assert r.returncode == 0, r.stdout + r.stderr
    assert "UnicodeEncodeError" not in r.stderr and "Traceback" not in r.stderr
