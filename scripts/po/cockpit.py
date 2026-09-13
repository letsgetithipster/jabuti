"""Cockpit xlsx gerado a partir de dados/ (spec §2, decisão 4): visão, não fonte.

Fonte de verdade continua sendo dados/ + política; o arquivo é regenerável e não
versionado. Única célula editável: 'Aporte do mês' (Aporte!B2). O resto é valor
gerado ou fórmula: Valor BRL = qty × cotação × câmbio; Blocos por SUMIFS; fila de
aporte por bloco = ordem por gap decrescente consumindo o aporte (fórmulas RANK/SUMIFS).
openpyxl é opcional: import lazy com erro acionável.

O número é o de po.carteira.valorar — o mesmo do ESTADO.md, nada recalculado aqui.
Os avisos da valoração (cotação ou câmbio velho, bloco sem banda) viajam com ele: vão
listados no fim do LEIAME e anunciados em Aporte!A1. Tela de aporte construída sobre
preço velho é exatamente o número plausível e errado que esta fase existe para caçar.
"""
import datetime
from pathlib import Path

from po.carteira import valorar
from po.config import caminho_planilhas, carregar_config
from po.ingestao.leitores import DependenciaAusente

MSG_OPENPYXL = "o cockpit xlsx exige openpyxl: pip install -r requirements-xlsx.txt"
COR_EDITAVEL = "FFF3D6"
COR_CABECALHO = "EDE7DB"
COR_AVISO = "9C2B00"
FMT_NUM = "#,##0.00"
FMT_PCT = "0.0%"


def gerar_cockpit(raiz: str | Path, agora: datetime.datetime | None = None) -> Path:
    """Gera <caminhos.planilhas>/cockpit.xlsx. ValueError se dados/ não sustenta o número."""
    try:
        import openpyxl
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise DependenciaAusente(MSG_OPENPYXL) from e
    raiz = Path(raiz)
    agora = agora or datetime.datetime.now()
    cfg = carregar_config(raiz)
    # a mesma data do carimbo do LEIAME decide o que é cotação velha: senão a planilha
    # diria "gerado em X" e mediria a idade dos preços contra outro dia.
    c = valorar(raiz, hoje=agora.date())
    destino = caminho_planilhas(raiz, cfg) / "cockpit.xlsx"
    destino.parent.mkdir(parents=True, exist_ok=True)

    negrito = Font(bold=True)
    fill_cab = PatternFill("solid", fgColor=COR_CABECALHO)
    fill_edit = PatternFill("solid", fgColor=COR_EDITAVEL)

    def cabecalho(ws, titulos, larguras):
        for i, (t, w) in enumerate(zip(titulos, larguras), start=1):
            cel = ws.cell(row=1, column=i, value=t)
            cel.font, cel.fill = negrito, fill_cab
            ws.column_dimensions[get_column_letter(i)].width = w
        ws.freeze_panes = "A2"

    wb = openpyxl.Workbook()
    # ---- LEIAME
    ws = wb.active
    ws.title = "LEIAME"
    ws.column_dimensions["A"].width = 110
    texto_leiame = [
        "jabuti — cockpit",
        f"Gerado em {agora:%d/%m/%Y %H:%M} por scripts/gerar_cockpit.py a partir de dados/ e da política declarada.",
        "Este arquivo é VISÃO, não fonte: regenerável, não versionado. A fonte de verdade é dados/ (CSVs) + politica/.",
        "Única célula editável: 'Aporte do mês' na aba Aporte (amarela). Tudo o mais é gerado ou fórmula; edição à mão se perde na próxima geração.",
        "Posições: Valor BRL = Qty × Cotação × Câmbio (fórmula). Cotação = última por data em cotacoes.csv; Câmbio = par {moeda}BRL mais recente (1 se BRL).",
        "'Result. moeda ativo' é o ganho na MOEDA DO ATIVO (Cotação ÷ PM − 1) e NÃO contém variação cambial: em posição USD, 8% aqui são 8% em dólar, não em reais.",
        "'Custo (câmbio hoje)' remarca o PM histórico pelo câmbio de HOJE, então não é o valor em reais que saiu da sua conta. Resultado em BRL exige o câmbio da data de cada compra, que esta fase ainda não guarda — por isso não existe coluna de resultado em BRL.",
        "Blocos: soma por classe (SUMIFS), % da carteira, banda declarada em politica/01-alocacao-alvo.md, desvio e gap até o alvo.",
        "Aporte: fila por BLOCO — ordem por gap decrescente após o aporte entrar; Sugerido consome o aporte na ordem. Fila por ticker chega na Fase 4 (exige tese validada).",
        "Convenção de hora: hora da fonte; fonte diária sem hora (PTAX, manual sem hora) grava 00:00.",
        "Isto executa a política que você declarou; não é recomendação de investimento.",
    ]
    if c.avisos:
        texto_leiame += ["", f"AVISOS ({len(c.avisos)}) — os números acima foram calculados apesar deles:"]
        texto_leiame += [f"- {a}" for a in c.avisos]
    for i, linha in enumerate(texto_leiame, start=1):
        cel = ws.cell(row=i, column=1, value=linha)
        cel.alignment = Alignment(wrap_text=True)
        if i == 1:
            cel.font = Font(bold=True, size=14)
        elif linha.startswith("AVISOS (") or linha.startswith("- "):
            cel.font = Font(bold=linha.startswith("AVISOS ("), color=COR_AVISO)

    # ---- Posições
    ws = wb.create_sheet("Posições")
    cabecalho(ws, ["Ticker", "Classe", "Conta", "Qty", "PM", "Moeda", "Cotação", "Data cot.", "Fonte", "Câmbio",
                   "Valor BRL", "Custo (câmbio hoje)", "Result. moeda ativo", "% carteira"],
              [10, 12, 14, 10, 12, 8, 12, 12, 10, 10, 14, 19, 19, 11])
    n = len(c.linhas)
    tot = n + 2
    for r, l in enumerate(c.linhas, start=2):
        for col, v in enumerate([l.ticker, l.classe, l.conta, l.qty, l.pm, l.moeda, l.preco, l.data_cotacao,
                                 l.fonte, l.cambio_cotacao], start=1):
            ws.cell(row=r, column=col, value=v)
        ws.cell(row=r, column=11, value=f"=D{r}*G{r}*J{r}").number_format = FMT_NUM
        ws.cell(row=r, column=12, value=round(l.custo_brl, 2)).number_format = FMT_NUM
        # Cotação/PM e não Valor/Custo: dá o mesmo número (os dois usam o câmbio de hoje) e
        # deixa visível na barra de fórmulas que o resultado é na moeda do ativo, sem câmbio.
        ws.cell(row=r, column=13, value=f'=IF(E{r}=0,"",G{r}/E{r}-1)').number_format = FMT_PCT
        ws.cell(row=r, column=14, value=f"=IF($K${tot}=0,0,K{r}/$K${tot})").number_format = FMT_PCT
        for col in (5, 7, 10):
            ws.cell(row=r, column=col).number_format = FMT_NUM
    ws.cell(row=tot, column=1, value="Total").font = negrito
    if n:
        ws.cell(row=tot, column=11, value=f"=SUM(K2:K{tot - 1})").number_format = FMT_NUM
        ws.cell(row=tot, column=12, value=f"=SUM(L2:L{tot - 1})").number_format = FMT_NUM
        ws.cell(row=tot, column=14, value=f"=SUM(N2:N{tot - 1})").number_format = FMT_PCT
    else:
        ws.cell(row=tot, column=11, value=0)
    fim_pos = max(tot - 1, 2)

    # ---- Blocos
    ws = wb.create_sheet("Blocos")
    cabecalho(ws, ["Bloco", "Valor BRL", "% carteira", "Mín %", "Alvo %", "Máx %", "Desvio", "Gap até o alvo"],
              [14, 14, 11, 8, 8, 8, 11, 15])
    bandas = {b.bloco: b for b in c.bandas}
    ordem = [b.bloco for b in c.bandas] + sorted(bl for bl in c.por_bloco if bl not in bandas)
    tb = len(ordem) + 2
    for r, bloco in enumerate(ordem, start=2):
        b = bandas.get(bloco)
        ws.cell(row=r, column=1, value=bloco)
        ws.cell(row=r, column=2, value=f"=SUMIFS(Posições!$K$2:$K${fim_pos},Posições!$B$2:$B${fim_pos},A{r})").number_format = FMT_NUM
        ws.cell(row=r, column=3, value=f"=IF($B${tb}=0,0,B{r}/$B${tb})").number_format = FMT_PCT
        if b:
            ws.cell(row=r, column=4, value=b.minimo)
            ws.cell(row=r, column=5, value=b.alvo)
            ws.cell(row=r, column=6, value=b.maximo)
        ws.cell(row=r, column=7, value=f'=IF(D{r}="","sem banda",IF(C{r}*100<D{r},"abaixo",IF(C{r}*100>F{r},"acima","dentro")))')
        ws.cell(row=r, column=8, value=f'=IF(E{r}="","",MAX(0,E{r}/100*$B${tb}-B{r}))').number_format = FMT_NUM
    ws.cell(row=tb, column=1, value="Total").font = negrito
    ws.cell(row=tb, column=2, value=f"=SUM(B2:B{tb - 1})" if ordem else 0).number_format = FMT_NUM
    ws.cell(row=tb, column=5, value=f"=SUM(E2:E{tb - 1})" if ordem else 0)

    # ---- Aporte
    ws = wb.create_sheet("Aporte")
    ws.column_dimensions["A"].width = 22
    for col in "BCDEF":
        ws.column_dimensions[col].width = 16
    if c.avisos:
        cel = ws.cell(row=1, column=1,
                      value=f"ATENÇÃO: {len(c.avisos)} aviso(s) sobre os dados desta planilha — "
                            "leia o fim da aba LEIAME antes de decidir o aporte.")
        cel.font = Font(bold=True, color=COR_AVISO)
    ws["A2"] = "Aporte do mês (R$)"
    ws["A2"].font = negrito
    ws["B2"] = 0
    ws["B2"].fill = fill_edit
    ws["B2"].number_format = FMT_NUM
    ws["C2"] = "← única célula editável"
    for i, t in enumerate(["Bloco", "Valor atual", "Alvo %", "Gap pós-aporte", "Ordem", "Sugerido"], start=1):
        cel = ws.cell(row=4, column=i, value=t)
        cel.font, cel.fill = negrito, fill_cab
    com_banda = [bl for bl in ordem if bl in bandas]
    ini, fim = 5, 5 + len(com_banda) - 1
    for r, bloco in enumerate(com_banda, start=ini):
        rb = ordem.index(bloco) + 2
        ws.cell(row=r, column=1, value=bloco)
        ws.cell(row=r, column=2, value=f"=Blocos!B{rb}").number_format = FMT_NUM
        ws.cell(row=r, column=3, value=f"=Blocos!E{rb}")
        ws.cell(row=r, column=4, value=f"=MAX(0,C{r}/100*(Blocos!$B${tb}+$B$2)-B{r})").number_format = FMT_NUM
        ws.cell(row=r, column=5, value=f'=IF(D{r}>0,RANK(D{r},$D${ini}:$D${fim},0)+COUNTIF($D${ini}:D{r},D{r})-1,"")')
        ws.cell(row=r, column=6, value=f'=IF(E{r}="","",MAX(0,MIN(D{r},$B$2-SUMIFS($D${ini}:$D${fim},$E${ini}:$E${fim},"<"&E{r}))))').number_format = FMT_NUM
    if com_banda:
        ws.cell(row=fim + 1, column=1, value="Total sugerido").font = negrito
        ws.cell(row=fim + 1, column=6, value=f"=SUM(F{ini}:F{fim})").number_format = FMT_NUM
        ws.cell(row=fim + 2, column=1, value="Sobra do aporte")
        ws.cell(row=fim + 2, column=6, value=f"=$B$2-F{fim + 1}").number_format = FMT_NUM
    else:
        ws.cell(row=ini, column=1, value="Sem bandas declaradas: defina em politica/01-alocacao-alvo.md (/jabuti-estrategia).")
    ws.freeze_panes = "A5"

    wb.save(destino)
    return destino
