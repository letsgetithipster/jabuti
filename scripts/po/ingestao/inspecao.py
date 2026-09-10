"""Dump determinístico de um documento de corretora, para a LLM desenhar o mapeamento
sem abrir o binário. Não converte nada: mostra as células como o arquivo guarda."""
import csv
from pathlib import Path

from po.ingestao.leitores import MSG_OPENPYXL, DependenciaAusente

DELIMITADORES = [",", ";", "\t", "|"]
N_PADRAO = 12


def _fmt(linha) -> str:
    return " | ".join("" if c is None else str(c).strip() for c in linha)


def inspecionar(caminho: str | Path, aba: str | int | None = None, n: int = N_PADRAO) -> str:
    caminho = Path(caminho)
    if not caminho.is_file():   # exists() é verdade para diretório
        return f"{caminho}: arquivo não encontrado"
    if caminho.suffix.lower() == ".xlsx":
        return _xlsx(caminho, aba, n)
    return _csv(caminho, n)


def _csv(caminho: Path, n: int) -> str:
    bruto = caminho.read_bytes()
    enc, texto = "latin-1", ""
    for candidato in ("utf-8-sig", "latin-1"):
        try:
            texto, enc = bruto.decode(candidato), candidato
            break
        except UnicodeDecodeError:
            continue
    amostra = texto.splitlines()[:20]
    delim = max(DELIMITADORES, key=lambda d: sum(l.count(d) for l in amostra))
    linhas = list(csv.reader(texto.splitlines(), delimiter=delim))
    out = [f"Arquivo: {caminho.name} · formato: csv · encoding: {enc} · delimitador: {delim!r} · {len(linhas)} linha(s)", ""]
    for i, l in enumerate(linhas[:n], start=1):
        out.append(f"{i:>4}: {_fmt(l)}")
    if len(linhas) > n:
        out.append(f"  … ({len(linhas) - n} linhas omitidas)")
        out.append(f"{len(linhas):>4}: {_fmt(linhas[-1])}")
    out += ["", "Bloco 'arquivo' sugerido para o mapeamento:",
            f"  arquivo: {{formato: csv, encoding: {enc}, delimitador: {delim!r}, cabecalho-contem: [...]}}"]
    return "\n".join(out)


def _xlsx(caminho: Path, aba, n: int) -> str:
    try:
        import openpyxl
    except ImportError as e:
        raise DependenciaAusente(MSG_OPENPYXL) from e
    wb = openpyxl.load_workbook(caminho, data_only=True)
    out = [f"Arquivo: {caminho.name} · formato: xlsx · abas: {wb.sheetnames}"]
    if aba is None:
        abas = wb.worksheets
    else:
        abas = [wb.worksheets[aba] if isinstance(aba, int) else wb[aba]]
    for ws in abas:
        linhas = [list(r) for r in ws.iter_rows(values_only=True)]
        out += ["", f"== aba {ws.title!r} ({len(linhas)} linhas)"]
        for i, l in enumerate(linhas[:n], start=1):
            marca = "   ← cabeçalho candidato" if sum(1 for c in l if isinstance(c, str) and c.strip()) >= 3 else ""
            out.append(f"{i:>4}: {_fmt(l)}{marca}")
        if len(linhas) > n:
            out.append(f"  … ({len(linhas) - n} linhas omitidas)")
            out.append(f"{len(linhas):>4}: {_fmt(linhas[-1])}")
    out += ["", "Datas e números aparecem como a célula guarda (datetime/float) ou como texto; "
                "o mapeamento declara os formatos em datas.formatos."]
    return "\n".join(out)
